import io
import json
import pandas as pd
import datetime
import unicodedata
import re
import os
import uuid
from PIL import Image, ImageOps
from openpyxl import load_workbook
from django.conf import settings
from rapidfuzz import fuzz
from django.db import connection, transaction
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from rest_framework import status
from accounts.views import login_required_api  # Importar el decorador de autenticación
from psycopg2.extras import execute_values


def normalize_key(key: str) -> str:
    return (
        key.strip()
        .lower()
        .replace(" ", "_")
        .replace(".", "_")  # Cambiar punto por guión bajo
        .replace("ó", "o")
        .replace("í", "i")
        .replace("ú", "u")
        .replace("é", "e")
        .replace("á", "a")
    )


# ================== FECHAS ==================
def _parse_fecha_valor(valor):
    """Intenta parsear un valor de fecha desde múltiples formatos:
    - Cadenas como 9/6/2019, 09-06-2019, 2019/06/09, etc.
    - Seriales de Excel (número de días desde 1899-12-30)
    Retorna pd.Timestamp o pd.NaT.
    """
    try:
        if valor is None:
            return pd.NaT

        # Si ya es fecha de pandas
        if isinstance(valor, (pd.Timestamp, datetime.date, datetime.datetime)):
            return pd.to_datetime(valor, errors="coerce")

        # Si es numérico, intentar como serial de Excel
        if isinstance(valor, (int, float)):
            # Rango razonable para serial de Excel
            if 1000 <= float(valor) <= 60000:
                base = pd.Timestamp("1899-12-30")
                try:
                    return base + pd.to_timedelta(int(valor), unit="D")
                except Exception:
                    pass
            # Si no, intentar como epoch seconds (menos probable)
            try:
                ts = pd.to_datetime(valor, unit="s", errors="coerce")
                if pd.notna(ts):
                    return ts
            except Exception:
                pass

        # Si es string, limpiar y probar múltiples estrategias
        s = str(valor).strip()
        if not s:
            return pd.NaT

        # Reemplazar separadores comunes
        s_norm = s.replace("\\", "/").replace("-", "/").replace(".", "/")

        # Intento 1: dayfirst=True (formato latino: D/M/Y)
        ts = pd.to_datetime(s_norm, dayfirst=True, errors="coerce")
        if pd.notna(ts):
            return ts

        # Intento 2: dayfirst=False (M/D/Y)
        ts = pd.to_datetime(s_norm, dayfirst=False, errors="coerce")
        if pd.notna(ts):
            return ts

        # Intento 3: formatos explícitos comunes
        formatos = [
            "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y", "%m/%d/%y",
            "%Y/%m/%d", "%Y-%m-%d",
        ]
        for fmt in formatos:
            try:
                return pd.to_datetime(datetime.datetime.strptime(s_norm, fmt))
            except Exception:
                continue

        return pd.NaT
    except Exception:
        return pd.NaT

def normalizar_columna_fecha_yyyy_mm_dd(serie: pd.Series) -> pd.Series:
    """Normaliza una Serie de fechas a formato YYYY-MM-DD con fallback a 2000-01-01."""
    # Aplicar parser robusto elemento a elemento
    fechas = serie.apply(_parse_fecha_valor)
    # Reemplazar NaT por fecha por defecto
    fechas = fechas.fillna(pd.Timestamp("2000-01-01"))
    # Asegurar tipo datetime y formatear
    fechas = pd.to_datetime(fechas, errors="coerce").dt.strftime("%Y-%m-%d")
    return fechas

# ================== FUNCIONES DE IMÁGENES ==================
# Función para extraer imágenes del Excel
def extraer_imagenes_excel(file_path, inventario_numero):
    """
    Extrae imágenes del Excel y las guarda como archivos físicos
    Retorna la ruta de la imagen si se encuentra una
    """
    try:
        # Cargar el workbook
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Buscar imágenes en el worksheet
        if hasattr(worksheet, '_images') and worksheet._images:
            # Si hay imágenes, tomar la primera (puedes modificar esta lógica)
            for img in worksheet._images:
                # Obtener los datos de la imagen
                image_data = img._data()
                
                # Generar nombre único para el archivo
                file_extension = 'png'  # Por defecto PNG
                if hasattr(img, 'format') and img.format:
                    file_extension = img.format.lower()
                
                filename = f"inventario_{inventario_numero}_{uuid.uuid4().hex[:8]}.{file_extension}"
                
                # Crear directorio si no existe
                images_dir = settings.MEDIA_ROOT / 'inventario_images'
                os.makedirs(images_dir, exist_ok=True)
                
                # Ruta completa del archivo
                file_path = images_dir / filename
                
                # Guardar la imagen
                with open(file_path, 'wb') as f:
                    f.write(image_data)
                
                # Retornar la ruta relativa para la base de datos
                return f"inventario_images/{filename}"
        
        return None
        
    except Exception as e:
        print(f"Error extrayendo imagen para inventario {inventario_numero}: {str(e)}")
        return None

# Función alternativa más robusta para extraer imágenes
def extraer_imagenes_excel_avanzado(file_path, inventario_numero, row_number=None):
    """
    Versión más avanzada que busca imágenes por fila específica
    """
    try:
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Buscar imágenes asociadas a una fila específica
        if hasattr(worksheet, '_images') and worksheet._images:
            for img in worksheet._images:
                # Verificar si la imagen está en la fila correcta
                if row_number and hasattr(img, 'anchor') and img.anchor:
                    # Si la imagen está cerca de la fila que estamos procesando
                    if abs(img.anchor.row - row_number) <= 1:  # Tolerancia de 1 fila
                        return procesar_imagen(img, inventario_numero)
                
                # Si no hay row_number específico, tomar la primera imagen
                if not row_number:
                    return procesar_imagen(img, inventario_numero)
        
        return None
        
    except Exception as e:
        print(f"Error extrayendo imagen avanzada para inventario {inventario_numero}: {str(e)}")
        return None

# ================== FUNCIONES DE ubicaciones ==================

# Cache de edificios con palabras clave
edificios_keywords = {
    "auditorio luis a. calvo": 1,
    "administracion 1": 2,
    "administracion 3": 3,
    "teatrino jose antonio galan": 4,
    "administracion 2": 5,
    "bienestar estudiantil": 6,
    "casona la perla": 7,
    "ciencias humanas 2": 8,
    "ingenieria mecanica": 9,
    "auditorio luis eduardo lobo c.": 10,
    "biblioteca": 11,
    "instituto de lenguas": 12,
    "ingenieria industrial": 13,
    "laboratorio biologia vegetal": 14,
    "laboratorios livianos": 15,
    "camilo torres": 16,
    "centic": 17,
    "uisalud": 18,
    "federico mamitza bayer": 19,
    "ingenieria e3t": 20,
    "facultad de ciencias": 21,
    "ingenieria quimica": 22,
    "aula maxima de ciencias": 23,
    "bienestar pro": 24,
    "laboratorio de alta tension": 25,
    "laboratorio de hidraulica": 26,
    "talleres de diseno industrial": 27,
    "planta de aceros": 28,
    "jorge bautista vesga": 29,
    "laboratorios pesados": 30,
    "daniel casas": 31,
    "edic": 32,
    "ciencias humanas (virginia gutierrez)": 33,
    "coliseo uis": 34,
    "cenivam": 35,
    "laboratorio estructura y materiales": 36,
    "gestion logistica": 37,
    "bienestar campestre": 38,
    "auditorio jorge zalama": 39,
}


def mapear_ubicacion(texto_ubicacion):
    if not texto_ubicacion:
        return None
    texto_norm = texto_ubicacion.strip().lower()
    
    for keyword, eid in edificios_keywords.items():
        if keyword in texto_norm:
            return eid
    return None


# ================== FUNCIONES DE IMÁGENES ==================

def procesar_imagen(img, inventario_numero):
    """
    Procesa una imagen individual y la guarda en formato WebP sin metadatos
    """
    try:
        # Obtener datos de la imagen
        image_data = img._data()
        
        # Crear objeto PIL Image desde los bytes
        with Image.open(io.BytesIO(image_data)) as original_image:
            
            # Aplicar corrección de orientación EXIF si existe
            original_image = ImageOps.exif_transpose(original_image)
            
            # Convertir a RGB si tiene canal alpha (PNG con transparencia)
            if original_image.mode in ('RGBA', 'LA'):
                background = Image.new('RGB', original_image.size, (255, 255, 255))
                mask = original_image.split()[-1] if original_image.mode in ('RGBA', 'LA') else None
                background.paste(original_image, mask=mask)
                cleaned_image = background
            elif original_image.mode == 'P':
                cleaned_image = original_image.convert('RGB')
            else:
                cleaned_image = original_image.convert('RGB')
            
            # Redimensionar si es muy grande
            if cleaned_image.width > 1920 or cleaned_image.height > 1080:
                cleaned_image.thumbnail((1920, 1080), Image.Resampling.LANCZOS)
            
            # Generar nombre único con extensión webp
            filename = f"inventario_{inventario_numero}_{uuid.uuid4().hex[:8]}.webp"
            
            # Crear directorio
            images_dir = settings.MEDIA_ROOT / 'inventario_images'
            os.makedirs(images_dir, exist_ok=True)
            
            # Ruta completa del archivo
            file_path = images_dir / filename
            
            # Guardar en formato WebP SIN METADATOS
            # La mejor forma de eliminar metadatos es recrear los datos de píxeles
            data = list(cleaned_image.getdata())
            new_image = Image.new(cleaned_image.mode, cleaned_image.size)
            new_image.putdata(data)
            
            # Guardar sin metadatos
            new_image.save(file_path, 'WEBP', quality=85, method=6, save_all=False)
            
            return f"inventario_images/{filename}"
        
    except Exception as e:
        print(f"Error procesando imagen: {str(e)}")
        return None

def mapear_imagenes_excel(file_path):
    """
    Mapea todas las imágenes del Excel por fila (1-based)
    """
    try:
        # Solo .xlsx soporta imágenes con openpyxl
        if not str(file_path).lower().endswith(".xlsx"):
            print("DEBUG: Archivo no es .xlsx; openpyxl no puede extraer imágenes de .xls")
            return {}

        wb = load_workbook(file_path)
        ws = wb.active

        imagen_map = {}

        if hasattr(ws, "_images") and ws._images:
            print(f"DEBUG: Imágenes en hoja: {len(ws._images)}")
            for i, img in enumerate(ws._images):
                anchor = getattr(img, "anchor", None)
                fila = None

                # Distintos tipos de anchor según openpyxl
                if anchor is not None:
                    # TwoCellAnchor / OneCellAnchor -> anchor._from.row (0-based)
                    base = getattr(anchor, "_from", None) or getattr(anchor, "from", None)
                    if base is not None and hasattr(base, "row"):
                        fila = int(base.row) + 1  # 1-based
                    elif hasattr(anchor, "row") and anchor.row is not None:
                        fila = int(anchor.row) + 1  # por si acaso

                if fila is not None:
                    imagen_map[fila] = img
                    print(f"DEBUG: Imagen {i} -> fila {fila}")
                else:
                    print(f"DEBUG: Imagen {i} sin fila detectable (anchor={anchor})")

        print(f"DEBUG: Filas con imagen: {sorted(imagen_map.keys())[:20]}")
        return imagen_map

    except Exception as e:
        print(f"Error mapeando imágenes: {e}")
        return {}
# ================== FUNCIONES DE NOMBRES ==================

def normalizar_texto(texto):
    if not texto:
        return ""
    texto = unicodedata.normalize("NFD", texto)
    texto = texto.encode("ascii", "ignore").decode("utf-8")
    texto = re.sub(r"[^A-Za-z\s]", "", texto).upper().strip()
    return texto

def dividir_nombre(nombre_completo):
    partes = normalizar_texto(nombre_completo).split()
    num = len(partes)

    if num == 0:
        return {"nombre1": "", "nombre2": "", "apellido1": "", "apellido2": ""}

    elif num == 1:
        return {"nombre1": partes[0], "nombre2": "", "apellido1": "", "apellido2": ""}

    elif num == 2:
        return {
            "nombre1": partes[0],
            "nombre2": "",
            "apellido1": partes[1],
            "apellido2": "",
        }

    elif num == 3:
        return {
            "nombre1": partes[0],
            "nombre2": "",
            "apellido1": partes[1],
            "apellido2": partes[2],
        }

    else:  # 4 o más
        return {
            "nombre1": partes[0],
            "nombre2": partes[1],
            "apellido1": partes[2],
            "apellido2": partes[3],
        }

def comparar_nombres_completos(
    n1,
    n2,
    umbral_global=85,
    umbrales_por_parte={"nombre1": 85, "nombre2": 75, "apellido1": 85, "apellido2": 75},
    doble_error_limite=2,
):
    p1 = dividir_nombre(n1)
    p2 = dividir_nombre(n2)

    scores = {
        "nombre1": fuzz.ratio(p1["nombre1"], p2["nombre1"]),
        "nombre2": fuzz.ratio(p1["nombre2"], p2["nombre2"]),
        "apellido1": fuzz.ratio(p1["apellido1"], p2["apellido1"]),
        "apellido2": fuzz.ratio(p1["apellido2"], p2["apellido2"]),
    }

    similitud_promedio = (
        0.25 * scores["nombre1"]
        + 0.15 * scores["nombre2"]
        + 0.3 * scores["apellido1"]
        + 0.3 * scores["apellido2"]
    )

    campos_en_riesgo = 0
    for parte, umbral in umbrales_por_parte.items():
        score = scores[parte]
        if score < umbral:
            return False, similitud_promedio, scores
        elif umbral <= score < 90:
            campos_en_riesgo += 1

    if campos_en_riesgo >= doble_error_limite:
        return False, similitud_promedio, scores

    return similitud_promedio >= umbral_global, similitud_promedio, scores


# ================== CATEGORÍAS ==================
CATEGORIA_MAP = {
    "Menores": 1,
    "Mayores": 2,
    "Intangible": 3,
}


# ================== CACHE INTELIGENTE DE USUARIOS ==================
class UsuarioCache:
    def __init__(self, cursor):
        self.cursor = cursor
        self.cache_directo = {}  # nombre -> id
        self.cache_invertido = {}  # nombre_invertido -> id
        self.cache_palabras = {}  # palabra -> lista de (id, nombre_completo)
        self.todos_usuarios = []  # lista de (id, nombre) para fuzzy
        self._construir_cache()
    
    def _construir_cache(self):
        """Construye todos los caches de una vez"""
        print("DEBUG: Construyendo cache de usuarios...")
        
        # Obtener todos los usuarios de una vez
        self.cursor.execute("SELECT id, nombre FROM usuarios")
        todos = self.cursor.fetchall()
        
        for user_id, nombre in todos:
            nombre_norm = nombre.strip()
            self.todos_usuarios.append((user_id, nombre_norm))
            
            # Cache directo
            self.cache_directo[nombre_norm.lower()] = user_id
            
            # Cache invertido
            invertido = " ".join(nombre_norm.split()[::-1])
            self.cache_invertido[invertido.lower()] = user_id
            
            # Cache por palabras
            palabras = nombre_norm.split()
            for palabra in palabras:
                if len(palabra) > 2:
                    palabra_key = palabra.lower()
                    if palabra_key not in self.cache_palabras:
                        self.cache_palabras[palabra_key] = []
                    self.cache_palabras[palabra_key].append((user_id, nombre_norm))
        
        print(f"DEBUG: Cache construido - {len(self.todos_usuarios)} usuarios")
    
    def buscar_usuario(self, nombre):
        """Búsqueda optimizada usando caches en memoria"""
        if not nombre or str(nombre).lower() == "nan":
            return None
        
        nombre_norm = nombre.strip()
        print(f"DEBUG: Buscando usuario: '{nombre_norm}'")
        
        # 1. Búsqueda directa
        user_id = self.cache_directo.get(nombre_norm.lower())
        if user_id:
            print(f"DEBUG: Usuario encontrado directo - ID: {user_id}")
            return user_id
        
        # 2. Búsqueda invertida
        invertido = " ".join(nombre_norm.split()[::-1])
        user_id = self.cache_invertido.get(invertido.lower())
        if user_id:
            print(f"DEBUG: Usuario encontrado invertido - ID: {user_id}")
            return user_id
        
        # 3. Búsqueda por palabras (LIKE)
        palabras = nombre_norm.split()
        candidatos = set()
        
        for palabra in palabras:
            if len(palabra) > 2:
                palabra_key = palabra.lower()
                if palabra_key in self.cache_palabras:
                    for user_id, nombre_db in self.cache_palabras[palabra_key]:
                        # Verificar que todas las palabras estén en el nombre de la DB
                        nombre_db_lower = nombre_db.lower()
                        if all(p.lower() in nombre_db_lower for p in palabras):
                            candidatos.add((user_id, nombre_db))
        
        if candidatos:
            # Tomar el primer candidato (más simple que fuzzy para LIKE)
            user_id, nombre_db = list(candidatos)[0]
            print(f"DEBUG: Usuario encontrado por palabras - ID: {user_id}, Nombre: '{nombre_db}'")
            return user_id
        
        # 4. Fuzzy matching (último recurso)
        mejor_match = None
        mejor_score = 0
        
        for user_id, nombre_db in self.todos_usuarios:
            valido, score, _ = comparar_nombres_completos(nombre_norm, nombre_db)
            if valido and score > mejor_score:
                mejor_match = user_id
                mejor_score = score
        
        if mejor_match:
            print(f"DEBUG: Usuario encontrado por FUZZY - ID: {mejor_match}, Score: {mejor_score}")
            return mejor_match
        
        print(f"DEBUG: Usuario NO encontrado para: '{nombre_norm}'")
        return 0


# ================== VALIDACION CATEGORÍA ==================
def validateCategory(nombreArchivo):
    if "Menores" in nombreArchivo:
        return CATEGORIA_MAP["Menores"]
    if "Mayores" in nombreArchivo:
        return CATEGORIA_MAP["Mayores"]
    else:
        return CATEGORIA_MAP["Intangible"]


# ================== IMPORTAR INVENTARIO ==================
@api_view(["POST"])
@parser_classes([MultiPartParser, FormParser])
@login_required_api
def importar_inventario(request):
    """
    Recibe un Excel/CSV y una categoría manual, limpia los datos y hace UPSERT en la tabla inventario_items (optimizado).
    Valida que el usuario autenticado sea el mismo que aparece en la columna 'recibido por'.
    Extrae imágenes del Excel y las guarda como archivos físicos.
    """
    file = request.FILES.get("file")
    categoria_manual = request.data.get("categoria")
    
    if not file:
        return Response(
            {"error": "Debes subir un archivo con key 'file'."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    
    # Validar que se envíe la categoría manualmente
    if not categoria_manual:
        return Response(
            {"error": "Debes enviar la categoría manualmente en el campo 'categoria'. Valores válidos: 'Menores', 'Mayores', 'Intangible'."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    
    # Validar que la categoría sea válida
    categoria_id = CATEGORIA_MAP.get(categoria_manual.strip())
    if not categoria_id:
        return Response(
            {"error": f"Categoría inválida: '{categoria_manual}'. Valores válidos: 'Menores', 'Mayores', 'Intangible'."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        # Guardar archivo temporalmente para extraer imágenes
        temp_file_path = None
        if file.name.endswith(('.xlsx', '.xls')):
            # Crear archivo temporal
            temp_file_path = f"/tmp/{uuid.uuid4().hex}_{file.name}"
            with open(temp_file_path, 'wb') as temp_file:
                for chunk in file.chunks():
                    temp_file.write(chunk)
            
            # Resetear el archivo para pandas
            file.seek(0)

        # --- Lectura del archivo ---
        if file.name.endswith(".xlsx"):
            read_file = pd.read_excel(file, engine="openpyxl")
            csv_buffer = io.StringIO()
            read_file.to_csv(csv_buffer, index=False)
            csv_buffer.seek(0)
            df = pd.read_csv(csv_buffer, skip_blank_lines=True, encoding="utf-8", header=9, skipinitialspace=True, usecols=lambda x: "Unnamed" not in x)
        elif file.name.endswith(".xls"):
            read_file = pd.read_excel(file, engine="xlrd")
            csv_buffer = io.StringIO()
            read_file.to_csv(csv_buffer, index=False)
            csv_buffer.seek(0)
            df = pd.read_csv(csv_buffer, skip_blank_lines=True, encoding="utf-8", header=9, skipinitialspace=True, usecols=lambda x: "Unnamed" not in x)
        else:
            df = pd.read_csv(file, skip_blank_lines=True, encoding="utf-8", header=9, skipinitialspace=True, usecols=lambda x: "Unnamed" not in x)

        # === DEBUG: Ver qué columnas tiene el archivo ===
        print(f"DEBUG: Total de columnas: {len(df.columns)}")
        print(f"DEBUG: Columnas encontradas: {list(df.columns)}")
        print(f"DEBUG: Primeras 3 filas:")
        print(df.head(3))

        # --- Usar la categoría manual ---
        df["Categoría"] = categoria_id

        # --- Limpieza ---
        try:
            # Buscar la columna de inventario con diferentes posibles nombres
            inventario_column = None
            possible_names = ["No. Inv.", "No Inv", "No. Inv", "Inventario", "No", "N°", "Nº", "No. Inventario"]
            
            for col_name in possible_names:
                if col_name in df.columns:
                    inventario_column = col_name
                    break
            
            if not inventario_column:
                return Response(
                    {"error": f"No se encontró la columna de inventario. Columnas disponibles: {list(df.columns)}"},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            
            print(f"DEBUG: Usando columna de inventario: '{inventario_column}'")
            
            # Renombrar la columna para normalización
            df = df.rename(columns={inventario_column: "No. Inv."})
            
            # Definir las columnas deseadas
            desired_columns = [
                "No. Inv.", "Descripción", "Marca", "Serial", "Valor", "Fecha Recibido", 
                "Ubicación", "Responsable", "Foto"
            ]
            
            # Verificar que todas las columnas necesarias existan
            missing_columns = []
            for col in desired_columns:
                if col not in df.columns:
                    missing_columns.append(col)
            
            if missing_columns:
                return Response(
                    {"error": f"Faltan las siguientes columnas en el archivo: {missing_columns}. Columnas disponibles: {list(df.columns)}"},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            
            # Limpiar datos
            df = df.dropna(subset=["No. Inv."])
            df = df[df["No. Inv."].astype(str).str.strip().str.isnumeric()]
            
            # Limpiar campos de texto
            for col in ["Marca", "Descripción", "Serial", "Responsable"]:
                if col in df.columns:
                    df[col] = df[col].astype(str).str.strip()
            
            if "Fecha Recibido" in df.columns:
                # Normalización robusta de fechas (soporta 9/6/2019, 09-06-2019, seriales Excel, etc.)
                df["Fecha Recibido"] = normalizar_columna_fecha_yyyy_mm_dd(df["Fecha Recibido"]) 

            
            if "Valor" in df.columns:
                df["Valor"] = df["Valor"].astype(str).str.replace(r"[^\d.]", "", regex=True)
                df["Valor"] = pd.to_numeric(df["Valor"], errors="coerce").fillna(0)
            
            # Seleccionar solo las columnas deseadas
            df = df[desired_columns]
            df.columns = [normalize_key(c) for c in df.columns]
            data_json = json.loads(df.to_json(orient="records", force_ascii=False))
            
            print(f"DEBUG: Datos procesados: {len(data_json)} registros")
            print(f"DEBUG: Primer registro: {data_json[0] if data_json else 'No hay datos'}")
            
        except Exception as e:
            return Response(
                {"error": f"Error procesando el archivo: {str(e)}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # === VALIDACIÓN DE USUARIO AUTENTICADO ===
        with connection.cursor() as cursor:
            # Obtener el nombre del usuario autenticado
            cursor.execute("SELECT nombre FROM usuarios WHERE id = %s", [request.user_id])
            usuario_autenticado = cursor.fetchone()
            
            if not usuario_autenticado:
                return Response(
                    {"error": "Usuario no encontrado en la base de datos."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            
            nombre_usuario_autenticado = usuario_autenticado[0]
            print(f"DEBUG: Usuario autenticado: '{nombre_usuario_autenticado}'")
            
            # Crear cache de usuarios para validación
            usuario_cache = UsuarioCache(cursor)
            
            # Validar que todos los registros tengan el mismo usuario en "responsable"
            usuarios_recibidos_ids = set()
            usuarios_recibidos_nombres = set()

            for item in data_json:
                responsable = item.get("responsable", "").strip()
                if responsable:
                    usuarios_recibidos_nombres.add(responsable)
                    usuario_id = usuario_cache.buscar_usuario(responsable)
                    if usuario_id and usuario_id != 0:
                        usuarios_recibidos_ids.add(usuario_id)

            # Verificar que solo haya un usuario único en "responsable"
            if len(usuarios_recibidos_ids) > 1:
                return Response(
                    {
                        "error": "Todos los registros deben tener el mismo usuario en la columna 'Responsable'.",
                        "usuarios_encontrados": list(usuarios_recibidos_nombres)
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # Verificar que el usuario en "responsable" coincida con el usuario autenticado
            if usuarios_recibidos_ids:
                usuario_id_en_archivo = list(usuarios_recibidos_ids)[0]
                
                if usuario_id_en_archivo != request.user_id:
                    cursor.execute("SELECT nombre FROM usuarios WHERE id = %s", [usuario_id_en_archivo])
                    usuario_encontrado = cursor.fetchone()
                    nombre_usuario_encontrado = usuario_encontrado[0] if usuario_encontrado else "Usuario desconocido"
                    
                    return Response(
                        {
                            "error": f"El usuario en el archivo ('{nombre_usuario_encontrado}') no coincide con tu usuario autenticado ('{nombre_usuario_autenticado}').",
                            "usuario_archivo": nombre_usuario_encontrado,
                            "usuario_autenticado": nombre_usuario_autenticado
                        },
                        status=status.HTTP_403_FORBIDDEN,
                    )

        not_found_ubicaciones, not_found_usuarios = [], []
        imagenes_procesadas = 0

        with transaction.atomic():
            with connection.cursor() as cursor:
                # === Cache de edificios ===
                cursor.execute("SELECT id, LOWER(edificio) FROM edificios")
                edificios_map = {nombre.strip().lower(): eid for eid, nombre in cursor.fetchall()}

                # === Cache inteligente de usuarios ===
                usuario_cache = UsuarioCache(cursor)

                # === MAPEAR IMÁGENES UNA SOLA VEZ (solo .xlsx) ===
                imagen_map = {}
                soporta_imagenes = bool(temp_file_path and str(temp_file_path).lower().endswith(".xlsx"))
                if soporta_imagenes:
                    imagen_map = mapear_imagenes_excel(temp_file_path)
                else:
                    print("DEBUG: Extracción de imágenes deshabilitada (requiere .xlsx)")

                # === Preparar registros para batch insert ===
               # Define el número de fila base según tu header (header=9 => primera data en fila 11)
                FILA_PRIMER_DATO = 11

                records = []
                inventarios_procesados = []
                for index, item in enumerate(data_json):
                    inventario_numero = item.get("no__inv_")
                    if not inventario_numero:
                        continue
                    inventarios_procesados.append(inventario_numero)
                    # === EXTRAER IMAGEN POR FILA (solo .xlsx) ===
                    foto = None
                    if soporta_imagenes and inventario_numero:
                        excel_row_number = FILA_PRIMER_DATO + index
                        img = imagen_map.get(excel_row_number)
                        if img:
                            foto = procesar_imagen(img, inventario_numero)
                            if foto:
                                imagenes_procesadas += 1
                                print(f"DEBUG: Inventario {inventario_numero} -> fila {excel_row_number} -> {foto}")
                        else:
                            # Opcional: tolerancia ±1 fila si el anclaje está pegado al borde
                            for delta in (-1, 1):
                                img_tol = imagen_map.get(excel_row_number + delta)
                                if img_tol:
                                    foto = procesar_imagen(img_tol, inventario_numero)
                                    if foto:
                                        imagenes_procesadas += 1
                                        print(f"DEBUG: Tolerancia: Inventario {inventario_numero} -> fila {excel_row_number+delta} -> {foto}")
                                    break
                    
                    # Ubicación
                    ubicacion_id = None
                    if item.get("ubicacion"):
                        ubicacion_id = mapear_ubicacion(item.get("ubicacion", ""))
                        if not ubicacion_id:
                            not_found_ubicaciones.append(item.get("ubicacion").strip())


                    # Usuario responsable (usar recibido_por_id)
                    responsable_id = usuario_cache.buscar_usuario(item.get("responsable", ""))

                    if responsable_id is None or responsable_id == 0:
                        not_found_usuarios.append(item.get("responsable"))
                        continue

                    # Agregar al batch (entregado_por = 1 por defecto, recibido_por = responsable_id)
                    records.append((
                        inventario_numero, item.get("descripcion"), item.get("marca"), 
                        item.get("serial"),
                        item.get("valor"), item.get("fecha_recibido"), categoria_id,
                        ubicacion_id, 1, responsable_id, 5,  # entregado_por = 1 (Luis Carlos), recibido_por = responsable_id
                        foto
                    ))
                                # === Consultar cuáles inventarios ya existen (antes de aplicar bloqueos) ===
                existentes = []
                nuevos = []

                if inventarios_procesados:
                    cursor.execute(
                        "SELECT inventario FROM inventario_items WHERE inventario = ANY(%s)",
                        (inventarios_procesados,),
                    )
                    existentes = [str(row[0]) for row in cursor.fetchall()]
                    nuevos = [str(inv) for inv in inventarios_procesados if str(inv) not in existentes]

                # === BLOQUEO: no importar items con trazabilidad COMPLETADO y accion baja/traslado ===
                elementos_completados = []
                inventarios_bloqueados = set()

                if inventarios_procesados:
                    # 1) Obtener ids de inventario_items para inventarios del archivo (solo los que ya existen)
                    cursor.execute(
                        """
                        SELECT id, inventario
                        FROM inventario_items
                        WHERE inventario = ANY(%s)
                        """,
                        (inventarios_procesados,),
                    )
                    rows_ids = cursor.fetchall()
                    id_to_inv = {rid: str(inv) for (rid, inv) in rows_ids}
                    ids_inventario_items = list(id_to_inv.keys())

                    # 2) Buscar trazas COMPLETADAS con accion baja/traslado
                    if ids_inventario_items:
                        cursor.execute(
                            """
                            SELECT
                                ii.id,
                                ii.inventario,
                                ii.descripcion,
                                ii.marca,
                                ii.serial,
                                ii.valor,
                                ii.fecha_recibido,
                                t.accion,
                                t.fecha AS fecha_movimiento,
                                t.detalle,
                                t.estado,
                                t.meta
                            FROM inventario_trazabilidad t
                            INNER JOIN inventario_items ii ON ii.id = t.inventario_id
                            WHERE t.inventario_id = ANY(%s)
                              AND LOWER(t.estado) = 'completado'
                              AND LOWER(t.accion) IN ('baja', 'traslado')
                            ORDER BY t.fecha DESC, t.id DESC
                            """,
                            (ids_inventario_items,),
                        )

                        cols = [c[0] for c in cursor.description]
                        for row in cursor.fetchall():
                            d = dict(zip(cols, row))

                            # marcar inventario como bloqueado
                            inventarios_bloqueados.add(str(d["inventario"]))

                            # serializar fechas
                            if d.get("fecha_recibido") and hasattr(d["fecha_recibido"], "isoformat"):
                                d["fecha_recibido"] = d["fecha_recibido"].isoformat()
                            if d.get("fecha_movimiento") and hasattr(d["fecha_movimiento"], "isoformat"):
                                d["fecha_movimiento"] = d["fecha_movimiento"].isoformat()

                            # parsear meta si viene como string
                            if d.get("meta") and isinstance(d["meta"], str):
                                try:
                                    d["meta"] = json.loads(d["meta"])
                                except Exception:
                                    pass

                            elementos_completados.append(d)

                # 3) Filtrar records para NO insertar/actualizar bloqueados
                if inventarios_bloqueados:
                    records = [rec for rec in records if str(rec[0]) not in inventarios_bloqueados]

                # 4) Recalcular inventarios_procesados para "procesados" (solo los que sí se importarán)
                inventarios_procesados_importados = [str(r[0]) for r in records]

                # 5) Recalcular existentes/nuevos respecto a los que sí se importarán
                existentes = []
                nuevos = []
                if inventarios_procesados_importados:
                    cursor.execute(
                        "SELECT inventario FROM inventario_items WHERE inventario = ANY(%s)",
                        (inventarios_procesados_importados,),
                    )
                    existentes = [str(row[0]) for row in cursor.fetchall()]
                    nuevos = [inv for inv in inventarios_procesados_importados if inv not in existentes]

                # 6) nuevos_detalle solo con los que sí se importan y son realmente nuevos
                nuevos_detalle = []
                if nuevos:
                    for rec in records:
                        inv_numero = str(rec[0])
                        if inv_numero in nuevos:
                            nuevos_detalle.append({
                                "inventario": rec[0],
                                "descripcion": rec[1],
                                "marca": rec[2],
                                "serial": rec[3],
                                "valor": rec[4],
                                "fecha_recibido": rec[5],
                                "categoria_id": rec[6],
                                "ubicacion_id": rec[7],
                                "entregado_por_id": rec[8],
                                "recibido_por_id": rec[9],
                                "escuela_id": rec[10],
                                "foto": rec[11],
                            })

                # 7) Batch UPSERT solo con los records filtrados
                if records:
                    execute_values(cursor, """
                        INSERT INTO inventario_items (
                            inventario, descripcion, marca, serial, valor, fecha_recibido,
                            categoria_id, ubicacion_id, entregado_por_id, recibido_por_id, escuela_id, foto
                        )
                        VALUES %s
                        ON CONFLICT (inventario) DO UPDATE SET
                            descripcion = EXCLUDED.descripcion,
                            marca = EXCLUDED.marca,
                            serial = EXCLUDED.serial,
                            valor = EXCLUDED.valor,
                            fecha_recibido = EXCLUDED.fecha_recibido,
                            categoria_id = EXCLUDED.categoria_id,
                            ubicacion_id = EXCLUDED.ubicacion_id,
                            entregado_por_id = EXCLUDED.entregado_por_id,
                            recibido_por_id = EXCLUDED.recibido_por_id,
                            escuela_id = EXCLUDED.escuela_id,
                            foto = EXCLUDED.foto
                        """,
                        records
                    )
                

                 # === Consultar elementos con estado "completado" en trazabilidad ===
                elementos_completados = []
                if inventarios_procesados:
                    # Primero obtener los IDs de inventario_items para los inventarios procesados
                    cursor.execute("""
                        SELECT id, inventario 
                        FROM inventario_items 
                        WHERE inventario = ANY(%s)
                    """, (inventarios_procesados,))
                    inventario_ids_map = {row[1]: row[0] for row in cursor.fetchall()}
                    
                    # Obtener los IDs de inventario_items que tienen estado "completado"
                    if inventario_ids_map:
                        inventario_ids_list = list(inventario_ids_map.values())
                        cursor.execute("""
                            SELECT DISTINCT
                                ii.id,
                                ii.inventario,
                                ii.descripcion,
                                ii.marca,
                                ii.serial,
                                ii.valor,
                                ii.fecha_recibido,
                                t.accion,
                                t.fecha as fecha_movimiento,
                                t.detalle,
                                t.estado,
                                t.meta
                            FROM inventario_items ii
                            INNER JOIN inventario_trazabilidad t ON ii.id = t.inventario_id
                            WHERE ii.id = ANY(%s)
                                AND LOWER(t.estado) = 'completado'
                            ORDER BY t.fecha DESC
                        """, (inventario_ids_list,))
                        
                        columns = [col[0] for col in cursor.description]
                        for row in cursor.fetchall():
                            elemento = dict(zip(columns, row))
                            # Convertir fecha a string si es necesario
                            if elemento.get('fecha_recibido'):
                                if hasattr(elemento['fecha_recibido'], 'isoformat'):
                                    elemento['fecha_recibido'] = elemento['fecha_recibido'].isoformat()
                            if elemento.get('fecha_movimiento'):
                                if hasattr(elemento['fecha_movimiento'], 'isoformat'):
                                    elemento['fecha_movimiento'] = elemento['fecha_movimiento'].isoformat()
                            # Parsear meta si es JSON string
                            if elemento.get('meta'):
                                try:
                                    elemento['meta'] = json.loads(elemento['meta']) if isinstance(elemento['meta'], str) else elemento['meta']
                                except:
                                    pass
                            elementos_completados.append(elemento)


    

        # Limpiar archivo temporal
        if temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)

        return Response(
            {
                "status": "ok",
                "procesados": len(records),
                "categoria_usada": categoria_manual,
                "imagenes_extraidas": imagenes_procesadas,
                "ubicaciones_no_encontradas": list(set(not_found_ubicaciones)),
                "usuarios_no_encontrados": list(set(not_found_usuarios)),
                "nuevos": len(nuevos),
                "repetidos": len(existentes),
                "inventarios_nuevos": nuevos_detalle,
                "elementos_completados": elementos_completados,
                "total_completados": len(elementos_completados),
                "bloqueados_por_completado": list(inventarios_bloqueados)
                
            },
            status=status.HTTP_201_CREATED,
        )

    except Exception as e:
        # Limpiar archivo temporal en caso de error
        if 'temp_file_path' in locals() and temp_file_path and os.path.exists(temp_file_path):
            os.remove(temp_file_path)
        return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

# ================== OBTENER INVENTARIO ==================
@api_view(["GET"])
@login_required_api
def obtener_inventario_usuario(request):
    try:
        user_id = request.user_id
        with connection.cursor() as cursor:
            cursor.execute("""
                SELECT 
                    ii.inventario, ii.descripcion, ii.marca, ii.serial, ii.valor, ii.fecha_recibido,
                    c.nombre as categoria, e.edificio as ubicacion,
                    ur.nombre as responsable, esc.nombre as escuela,
                    ii.foto
                FROM inventario_items ii
                LEFT JOIN categorias c ON ii.categoria_id = c.id
                LEFT JOIN edificios e ON ii.ubicacion_id = e.id
                LEFT JOIN usuarios ur ON ii.recibido_por_id = ur.id
                LEFT JOIN escuelas esc ON ii.escuela_id = esc.id
                WHERE ii.recibido_por_id = %s
                ORDER BY ii.fecha_recibido DESC
            """, [user_id])

            columns = [col[0] for col in cursor.description]
            items = []
            for row in cursor.fetchall():
                item = dict(zip(columns, row))
                # Construir URL completa de la imagen si existe
                if item.get('foto'):
                    base_url = request.build_absolute_uri('/')[:-1]
                    item['imagen_url'] = f"{base_url}{settings.MEDIA_URL}{item['foto']}"
                else:
                    item['imagen_url'] = None
                items.append(item)

        return Response({"success": True, "total_items": len(items), "items": items}, status=status.HTTP_200_OK)
    except Exception as e:
        return Response({"error": f"Error al obtener inventario: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)